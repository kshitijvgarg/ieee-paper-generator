from __future__ import annotations

import re
from copy import deepcopy
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Cm
from docx.text.paragraph import Paragraph

from .markdown_parser import Run, parse_body_blocks, parse_runs
from .spec import Author, Cell, Figure, PaperSpec, Section, Table
from .strict_to_transitional import ensure_transitional

DEFAULT_TEMPLATE = (
    Path(__file__).resolve().parent.parent / "templates" / "docx" / "ieee-conference-template-a4.docx"
)

# Map styleId (what appears in the template's XML) to human style name (what
# python-docx expects as a lookup key). See styles.xml in the template.
_STYLE_NAMES: dict[str, str] = {
    "papertitle": "paper title",
    "Author": "Author",
    "Abstract": "Abstract",
    "Keywords": "Keywords",
    "Heading1": "heading 1",
    "Heading2": "heading 2",
    "Heading3": "heading 3",
    "Heading4": "heading 4",
    "Heading5": "heading 5",
    "BodyText": "Body Text",
    "figurecaption": "figure caption",
    "references": "references",
    "bulletlist": "bullet list",
    "equation": "equation",
    "tablehead": "table head",
    "tablecolhead": "table col head",
    "tablecolsubhead": "table col subhead",
    "tablecopy": "table copy",
    "tablefootnote": "table footnote",
    "sponsors": "sponsors",
}

def _author_row_layout(n: int) -> list[tuple[int, int, int]]:
    """Decide how to lay out `n` authors across rows.

    Returns a list of (col_count, start_idx, end_idx) per row. Authors
    [start_idx:end_idx] of the input list belong to a row rendered in
    `col_count` columns. Matches common IEEE conventions:
        1 → 1×1,   2 → 2×1,   3 → 3×1,
        4 → 2×2,   5 → 3+2,   6 → 3×2.
    """
    if n <= 0 or n > 6:
        raise ValueError(f"Expected 1-6 authors, got {n}")
    if n <= 3:
        return [(n, 0, n)]
    if n == 4:
        return [(2, 0, 2), (2, 2, 4)]
    if n == 5:
        return [(3, 0, 3), (2, 3, 5)]
    return [(3, 0, 3), (3, 3, 6)]


def build_paper(
    spec: PaperSpec,
    output_path: str | Path,
    template_path: str | Path = DEFAULT_TEMPLATE,
) -> None:
    """Render `spec` into a new .docx using the IEEE template as a base."""
    template_path = Path(template_path)
    readable_template = ensure_transitional(template_path)
    doc = Document(str(readable_template))
    _set_title(doc, spec.title)
    _remove_subtitle_note(doc)
    _clear_body_after_authors(doc)
    _set_authors(doc, spec.authors)

    # Paper-wide counters passed by reference through a mutable dict so each
    # section can auto-number equations, figures, and tables in document
    # order.
    ctx: dict[str, int] = {"equation": 0, "figure": 0, "table": 0}

    if spec.sponsors:
        _add_sponsors(doc, spec.sponsors)

    _add_abstract(doc, spec.abstract)
    _add_keywords(doc, spec.keywords)

    for section in spec.sections:
        _add_section(doc, section, level=1, ctx=ctx, base_dir=spec.base_dir)

    # Paper-level figures (if any) render after all sections.
    for fig in spec.figures:
        ctx["figure"] += 1
        _add_figure(doc, fig, ctx["figure"], spec.base_dir)

    if spec.references:
        _add_references(
            doc,
            spec.references,
            numbered=spec.references_numbered,
            body=spec.references_body,
        )

    if spec.trailing_text.strip():
        _add_trailing_text(doc, spec.trailing_text)

    _finalise_body_sectPr(doc)
    doc.save(str(output_path))


def _add_trailing_text(doc: DocumentType, text: str) -> None:
    for kind, chunk in parse_body_blocks(text):
        style_key = "bulletlist" if kind == "bullet" else "BodyText"
        p = doc.add_paragraph(style=_STYLE_NAMES[style_key])
        _append_runs(p, parse_runs(chunk))


# --- Title ---------------------------------------------------------------


def _set_title(doc: DocumentType, title: str) -> None:
    for p in doc.paragraphs:
        if _style_id_of(p) == "papertitle":
            _replace_paragraph_text(p, parse_runs(title))
            return
    raise RuntimeError("Template is missing a 'papertitle' paragraph")


def _remove_subtitle_note(doc: DocumentType) -> None:
    """Delete the template's "Sub-titles are not captured in Xplore" paragraph.

    The template ships it as an Author-styled paragraph wedged between the
    title and the author grid. It's authoring guidance for the IEEE template
    user, not content we want in generated papers, so strip it."""
    body = doc.element.body
    for p in list(body):
        if p.tag != qn("w:p"):
            continue
        texts = p.findall(".//" + qn("w:t"))
        joined = "".join(t.text or "" for t in texts)
        if "Sub-titles are not captured" in joined:
            body.remove(p)
            return


# --- Authors -------------------------------------------------------------


def _set_authors(doc: DocumentType, authors: list[Author]) -> None:
    """Dynamically rebuild the template's author section based on the number
    of authors.

    The template hard-codes a 6-slot, 3-column × 2-row grid. Feeding fewer
    (or differently-grouped) authors requires changing both the paragraph
    content *and* the sectPrs that define each row's column count. We
    replace both the 4 Author grid paragraphs and the 2 spacer paragraphs
    that hold the 3-col sectPrs with freshly generated paragraphs sized to
    the author count.
    """
    rows = _author_row_layout(len(authors))

    body = doc.element.body
    all_ps = [c for c in body if c.tag == qn("w:p")]

    author_ps = [p for p in all_ps if _style_id_of_elem(p) == "Author"]
    # Template ships the 6-slot author grid as the last 4 Author-styled
    # paragraphs (2 spacer paragraphs precede them). We take the trailing 4
    # so this stays robust to the sub-title note's presence or removal.
    if len(author_ps) < 4:
        raise RuntimeError(
            f"Template missing expected Author paragraphs (found {len(author_ps)}, need >=4)"
        )
    grid_paras = author_ps[-4:]

    three_col_sectbreaks: list = []
    for p in all_ps:
        pPr = p.find(qn("w:pPr"))
        if pPr is None:
            continue
        sectPr = pPr.find(qn("w:sectPr"))
        if sectPr is None:
            continue
        cols = sectPr.find(qn("w:cols"))
        if cols is not None and cols.get(qn("w:num")) == "3":
            three_col_sectbreaks.append(p)
    if not three_col_sectbreaks:
        raise RuntimeError("Template is missing 3-column author sectPrs")

    template_author_sectPr = deepcopy(
        three_col_sectbreaks[0].find(qn("w:pPr")).find(qn("w:sectPr"))
    )

    first_grid = grid_paras[0]
    for col_count, start, end in rows:
        row_authors = authors[start:end]
        row_sectPr = _clone_sectPr_with_cols(template_author_sectPr, col_count)
        # Emit the row as two paragraphs: the content paragraph (cells with
        # column breaks, no sectPr) followed by an empty spacer paragraph
        # whose pPr carries the sectPr. This mirrors the template's
        # structure (content paragraphs then a trailing sectPr-bearing
        # spacer) and preserves the visual whitespace between author rows
        # and the abstract.
        first_grid.addprevious(_build_author_row_paragraph(row_authors))
        first_grid.addprevious(_build_sectbreak_spacer(row_sectPr))

    for old in grid_paras + three_col_sectbreaks:
        body.remove(old)


def _build_sectbreak_spacer(sectPr) -> "object":
    """Empty Normal-styled paragraph whose pPr holds `sectPr` and ends the
    author row's column-count section. The visible gap between authors and
    abstract is added separately via `spacing.before` on the abstract — more
    reliable across renderers than relying on an empty paragraph's height."""
    p = OxmlElement("w:p")
    pPr = OxmlElement("w:pPr")
    pPr.append(sectPr)
    p.append(pPr)
    return p


def _clone_sectPr_with_cols(template_sectPr, col_count: int):
    """Deep-copy a template sectPr and adjust its `<w:cols>` to `col_count`."""
    new = deepcopy(template_sectPr)
    cols = new.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        new.append(cols)
    if col_count <= 1:
        cols.attrib.pop(qn("w:num"), None)
    else:
        cols.set(qn("w:num"), str(col_count))
    if cols.get(qn("w:space")) is None:
        cols.set(qn("w:space"), "36pt")
    return new


def _build_author_row_paragraph(row_authors: list[Author]) -> "object":
    """Create a new Author-styled paragraph containing one cell per author in
    `row_authors`, separated by column breaks. The sectPr defining the row's
    column count lives on a separate spacer paragraph emitted by
    `_build_sectbreak_spacer`."""
    p = OxmlElement("w:p")
    pPr = OxmlElement("w:pPr")
    pStyle = OxmlElement("w:pStyle")
    pStyle.set(qn("w:val"), "Author")
    pPr.append(pStyle)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:before"), "5pt")
    spacing.set(qn("w:beforeAutospacing"), "1")
    pPr.append(spacing)
    p.append(pPr)

    for cell_idx, author in enumerate(row_authors):
        if cell_idx > 0:
            p.append(_make_column_break_run())
        lines = [
            (author.name, False),
            (author.dept, True),
            (author.org, True),
            (_city_country(author), False),
            (author.email, False),
        ]
        first_in_cell = True
        for text, italic in lines:
            for part in text.split("\n"):
                if not first_in_cell:
                    p.append(_make_line_break_run())
                p.append(_make_text_run(part, italic=italic, size_half_pt=18))
                first_in_cell = False
    return p


def _city_country(a: Author) -> str:
    parts = [a.city.strip(), a.country.strip()]
    return ", ".join(p for p in parts if p)


def _make_text_run(text: str, italic: bool = False, size_half_pt: int | None = None) -> OxmlElement:
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    if italic:
        rpr.append(OxmlElement("w:i"))
    if size_half_pt is not None:
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(size_half_pt))
        rpr.append(sz)
        szcs = OxmlElement("w:szCs")
        szcs.set(qn("w:val"), str(size_half_pt))
        rpr.append(szcs)
    if len(rpr):
        r.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    if text != text.strip():
        t.set(qn("xml:space"), "preserve")
    r.append(t)
    return r


def _make_line_break_run() -> OxmlElement:
    r = OxmlElement("w:r")
    br = OxmlElement("w:br")
    r.append(br)
    return r


def _make_column_break_run() -> OxmlElement:
    r = OxmlElement("w:r")
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "column")
    r.append(br)
    return r


# --- Clear body ----------------------------------------------------------


def _clear_body_after_authors(doc: DocumentType) -> None:
    """Delete template body content, preserving author-section breaks and the
    body-section properties.

    Template structure (by column count declared in sectPr):
        1-col title section   (sectPr in a pPr near the top)
        3-col author row 1    (sectPr in empty spacer paragraph after authors)
        3-col author row 2    (sectPr in next empty spacer paragraph)
        2-col body            (sectPr in the LAST body paragraph)
        terminal sectPr       (direct child of <w:body>)

    The paragraphs that *hold* the 3-col sectPrs are unstyled spacers after the
    last Author-styled paragraph, so we cannot just delete everything after the
    Author paragraphs. Instead we keep every paragraph up through the last
    3-col sectPr paragraph, drop everything in between, and stash the 2-col
    sectPr so we can re-attach it to the last new body paragraph in
    `_finalise_body_sectPr`.
    """
    body = doc.element.body
    paragraphs = [c for c in body if c.tag == qn("w:p")]

    preserved_body_sectPr = None
    author_boundary_idx = None
    for idx, p in enumerate(paragraphs):
        pPr = p.find(qn("w:pPr"))
        if pPr is None:
            continue
        sectPr = pPr.find(qn("w:sectPr"))
        if sectPr is None:
            continue
        cols = sectPr.find(qn("w:cols"))
        num_cols = cols.get(qn("w:num"), "1") if cols is not None else "1"
        if num_cols == "3":
            author_boundary_idx = idx  # track the latest 3-col paragraph
        elif num_cols == "2":
            preserved_body_sectPr = deepcopy(sectPr)

    if author_boundary_idx is None:
        raise RuntimeError("Template is missing the 3-column author sectPrs")
    if preserved_body_sectPr is None:
        raise RuntimeError("Template is missing the 2-column body sectPr")

    author_boundary = paragraphs[author_boundary_idx]
    body_children = list(body)
    start_remove = body_children.index(author_boundary) + 1
    for child in body_children[start_remove:]:
        if child.tag == qn("w:sectPr"):
            continue  # keep the terminal body sectPr
        body.remove(child)

    # Park the preserved sectPr on the document object; we re-attach in
    # `_finalise_body_sectPr` after all new content is appended.
    setattr(doc, "_ieee_body_sectPr", preserved_body_sectPr)


def _finalise_body_sectPr(doc: DocumentType) -> None:
    """Re-attach the stashed 2-col sectPr to the last body paragraph so the
    body section retains its two-column layout."""
    preserved = getattr(doc, "_ieee_body_sectPr", None)
    if preserved is None:
        return
    body = doc.element.body
    last_p = None
    for child in body:
        if child.tag == qn("w:p"):
            last_p = child
    if last_p is None:
        return
    pPr = last_p.find(qn("w:pPr"))
    if pPr is None:
        pPr = OxmlElement("w:pPr")
        last_p.insert(0, pPr)
    existing = pPr.find(qn("w:sectPr"))
    if existing is not None:
        pPr.remove(existing)
    pPr.append(preserved)


# --- Abstract & keywords -------------------------------------------------


def _add_abstract(doc: DocumentType, text: str) -> None:
    if not text.strip():
        return
    normalised = re.sub(r"\s+", " ", text).strip()
    p = doc.add_paragraph(style=_STYLE_NAMES["Abstract"])
    # Force a visible vertical gap between the author block and the abstract.
    # The empty spacer paragraphs that carry the row sectPrs render at near
    # zero height in LibreOffice, so we recreate the template's visible line
    # of whitespace by asking for 12pt of space-before on this paragraph.
    _set_spacing_before(p, "12pt")
    _append_runs(
        p,
        [
            Run("Abstract\u2014", bold=True, italic=True),
            *parse_runs(normalised),
        ],
        default_italic=True,
    )


def _set_spacing_before(paragraph: Paragraph, value: str) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    existing = pPr.find(qn("w:spacing"))
    if existing is None:
        existing = OxmlElement("w:spacing")
        pPr.append(existing)
    existing.set(qn("w:before"), value)
    # Disable Word's auto-spacing, which would otherwise override our value.
    existing.set(qn("w:beforeAutospacing"), "0")


def _add_keywords(doc: DocumentType, keywords: list[str]) -> None:
    if not keywords:
        return
    p = doc.add_paragraph(style=_STYLE_NAMES["Keywords"])
    joined = ", ".join(k.strip() for k in keywords if k.strip())
    _append_runs(
        p,
        [
            Run("Keywords\u2014", bold=True, italic=True),
            Run(joined),
        ],
        default_italic=True,
    )


# --- Sections ------------------------------------------------------------


def _add_section(
    doc: DocumentType,
    section: Section,
    level: int,
    ctx: dict[str, int],
    base_dir: Path,
) -> None:
    heading_key = "Heading5" if not section.numbered else f"Heading{min(level, 5)}"
    heading_p = doc.add_paragraph(style=_STYLE_NAMES[heading_key])
    _append_runs(heading_p, parse_runs(section.heading))

    for kind, text in parse_body_blocks(section.body):
        if kind == "equation":
            ctx["equation"] += 1
            _add_equation(doc, text, ctx["equation"])
            continue
        style_key = "bulletlist" if kind == "bullet" else "BodyText"
        body_p = doc.add_paragraph(style=_STYLE_NAMES[style_key])
        _append_runs(body_p, parse_runs(text))

    for table in section.tables:
        ctx["table"] += 1
        _add_table(doc, table, ctx["table"])

    for fig in section.figures:
        ctx["figure"] += 1
        _add_figure(doc, fig, ctx["figure"], base_dir)

    for sub in section.subsections:
        _add_section(doc, sub, level + 1, ctx, base_dir)


# --- Figures -------------------------------------------------------------


def _add_figure(doc: DocumentType, fig: Figure, number: int, base_dir: Path) -> None:
    if fig.path:
        img_path = Path(fig.path)
        if not img_path.is_absolute():
            img_path = (base_dir / img_path).resolve()
        if not img_path.is_file():
            raise FileNotFoundError(f"Figure image not found: {img_path}")
        doc.add_picture(str(img_path), width=Cm(8.4))  # ≈ single column width on A4
    # `figure caption` style auto-generates "Fig. N." via numbering (numId 2),
    # so we emit only the caption text and the counter `number` is unused
    # here — retained for API symmetry.
    del number  # silence unused-arg lints
    caption_p = doc.add_paragraph(style=_STYLE_NAMES["figurecaption"])
    _append_runs(caption_p, parse_runs(fig.caption))


# --- Equations -----------------------------------------------------------


def _add_equation(doc: DocumentType, text: str, number: int) -> None:
    """Render an equation paragraph in the template's `equation` style.

    The style defines a centering tab and a right tab for the equation
    number. Italic variables in the source can be marked with `*x*`.
    """
    p = doc.add_paragraph(style=_STYLE_NAMES["equation"])
    p_elem = p._p

    lead_tab = OxmlElement("w:r")
    lead_tab.append(OxmlElement("w:tab"))
    p_elem.append(lead_tab)

    for run_spec in parse_runs(text):
        r = p.add_run(run_spec.text)
        if run_spec.italic:
            r.italic = True
        if run_spec.bold:
            r.bold = True

    trail_tab = OxmlElement("w:r")
    trail_tab.append(OxmlElement("w:tab"))
    p_elem.append(trail_tab)
    p.add_run(f"({number})")


# --- Tables --------------------------------------------------------------


def _add_table(doc: DocumentType, table_spec: "Table", number: int) -> None:
    """Render a bordered, centred table preceded by a "Table N. Caption"
    heading in the `table head` style.

    Supports cell merging via `rowspan`/`colspan` on individual cells. Row 0
    cells default to `table col head`, row 1 to `table col subhead` if
    `has_subhead_row` is set, otherwise `table copy`. Per-cell `style` on a
    `Cell` overrides the row default. Footnote (if any) renders below in the
    `table footnote` style.
    """
    # `table head` style auto-generates "Table I." via numbering (numId 9),
    # so we emit only the caption text. `number` kept for API symmetry.
    del number
    head = doc.add_paragraph(style=_STYLE_NAMES["tablehead"])
    _append_runs(head, parse_runs(table_spec.caption))

    if not table_spec.rows:
        return

    placements = _compute_cell_placements(table_spec.rows)
    nrows = len(table_spec.rows)
    ncols = max(c + cs for (_, c, _, cs, _) in placements)

    t = doc.add_table(rows=nrows, cols=ncols)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    _apply_table_borders(t)
    _size_table_to_column(t, ncols)

    for r, c, rowspan, colspan, cell in placements:
        target = t.cell(r, c)
        if rowspan > 1 or colspan > 1:
            target.merge(t.cell(r + rowspan - 1, c + colspan - 1))
        style_key = cell.style or _default_row_style_key(r, table_spec)
        cell_p = target.paragraphs[0]
        cell_p.style = doc.styles[_STYLE_NAMES[style_key]]
        _append_runs(cell_p, parse_runs(cell.text))

    if table_spec.footnote:
        fn = doc.add_paragraph(style=_STYLE_NAMES["tablefootnote"])
        _append_runs(fn, parse_runs(table_spec.footnote))


def _compute_cell_placements(
    rows: list[list[Cell]],
) -> list[tuple[int, int, int, int, Cell]]:
    """Walk user-supplied rows top-to-bottom, left-to-right and place each
    Cell at the next open (row, col) position, skipping positions already
    reserved by a prior cell's rowspan or colspan.

    Returns a list of (row, col, rowspan, colspan, cell) tuples in placement
    order. The caller only emits cells at these positions; covered positions
    are implicit in the merge.
    """
    reserved: set[tuple[int, int]] = set()
    placements: list[tuple[int, int, int, int, Cell]] = []
    for r, row in enumerate(rows):
        c = 0
        for cell in row:
            while (r, c) in reserved:
                c += 1
            placements.append((r, c, cell.rowspan, cell.colspan, cell))
            for dr in range(cell.rowspan):
                for dc in range(cell.colspan):
                    if dr == 0 and dc == 0:
                        continue
                    reserved.add((r + dr, c + dc))
            c += cell.colspan
    return placements


def _default_row_style_key(row_idx: int, table_spec: "Table") -> str:
    if row_idx == 0:
        return "tablecolhead"
    if row_idx == 1 and table_spec.has_subhead_row:
        return "tablecolsubhead"
    return "tablecopy"


def _apply_table_borders(table) -> None:
    tbl = table._tbl
    tblPr = tbl.find(qn("w:tblPr"))
    if tblPr is None:
        return
    existing = tblPr.find(qn("w:tblBorders"))
    if existing is not None:
        tblPr.remove(existing)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{edge}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "2")
        b.set(qn("w:space"), "0")
        b.set(qn("w:color"), "auto")
        borders.append(b)
    tblPr.append(borders)


def _size_table_to_column(table, ncols: int) -> None:
    """Constrain the table to fit a single body column on A4 (~240pt wide).
    python-docx otherwise lays tables out at full page width, which overflows
    the 2-column body section and clips cells."""
    # 1pt == 20 twips. 240pt column ≈ 4800 twips. Leave a little slack.
    total_twips = 4500
    per_col = total_twips // ncols

    tbl = table._tbl
    tblPr = tbl.find(qn("w:tblPr"))
    if tblPr is not None:
        for existing in list(tblPr.findall(qn("w:tblW"))) + list(tblPr.findall(qn("w:tblLayout"))):
            tblPr.remove(existing)
        tblW = OxmlElement("w:tblW")
        tblW.set(qn("w:w"), str(total_twips))
        tblW.set(qn("w:type"), "dxa")
        tblPr.append(tblW)
        layout = OxmlElement("w:tblLayout")
        layout.set(qn("w:type"), "fixed")
        tblPr.append(layout)

    for col in tbl.findall(qn("w:tblGrid") + "/" + qn("w:gridCol")):
        col.set(qn("w:w"), str(per_col))

    for row in tbl.findall(qn("w:tr")):
        for cell in row.findall(qn("w:tc")):
            tcPr = cell.find(qn("w:tcPr"))
            if tcPr is None:
                continue
            for existing in tcPr.findall(qn("w:tcW")):
                tcPr.remove(existing)
            tcW = OxmlElement("w:tcW")
            tcW.set(qn("w:w"), str(per_col))
            tcW.set(qn("w:type"), "dxa")
            tcPr.append(tcW)


def _roman(n: int) -> str:
    vals = [(1000, "M"), (900, "CM"), (500, "D"), (400, "CD"),
            (100, "C"), (90, "XC"), (50, "L"), (40, "XL"),
            (10, "X"), (9, "IX"), (5, "V"), (4, "IV"), (1, "I")]
    out = []
    for v, s in vals:
        while n >= v:
            out.append(s)
            n -= v
    return "".join(out)


# --- Sponsor footnote frame ----------------------------------------------


def _add_sponsors(doc: DocumentType, text: str) -> None:
    """Insert a framed paragraph in the `sponsors` style anchored to the
    bottom-left of page 1, matching the template's sponsor footnote box."""
    body = doc.element.body
    title_p = None
    for child in body:
        if child.tag == qn("w:p"):
            title_p = child
            break
    if title_p is None:
        return

    p = OxmlElement("w:p")
    pPr = OxmlElement("w:pPr")
    pStyle = OxmlElement("w:pStyle")
    pStyle.set(qn("w:val"), "sponsors")
    pPr.append(pStyle)
    framePr = OxmlElement("w:framePr")
    framePr.set(qn("w:wrap"), "auto")
    framePr.set(qn("w:vAnchor"), "page")
    framePr.set(qn("w:hAnchor"), "page")
    framePr.set(qn("w:x"), "45.90pt")
    framePr.set(qn("w:y"), "756.05pt")
    pPr.append(framePr)
    p.append(pPr)

    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    r.append(t)
    p.append(r)

    title_p.addnext(p)


# --- References ----------------------------------------------------------


_LEADING_REF_NUM = re.compile(r"^\s*\[\d+\]\s*")


def _add_references(
    doc: DocumentType,
    refs: list[str],
    numbered: bool,
    body: str = "",
) -> None:
    # IEEE convention: References heading is un-numbered, rendered via
    # Heading 5 (which has no numPr in the template's styles.xml). If the
    # caller explicitly opts in to a numbered heading, use Heading 1.
    heading_key = "Heading1" if numbered else "Heading5"
    heading = doc.add_paragraph(style=_STYLE_NAMES[heading_key])
    _append_runs(heading, [Run("References")])
    # Optional instructional body (IEEE template ships authoring guidance
    # under the References heading before the numbered list).
    if body.strip():
        for kind, text in parse_body_blocks(body):
            style_key = "bulletlist" if kind == "bullet" else "BodyText"
            body_p = doc.add_paragraph(style=_STYLE_NAMES[style_key])
            _append_runs(body_p, parse_runs(text))
    for ref in refs:
        # The `references` style auto-generates [1], [2], ... via numbering.
        # Users often write their refs with an explicit "[1] " prefix; strip
        # it so we don't double-number.
        cleaned = _LEADING_REF_NUM.sub("", ref)
        p = doc.add_paragraph(style=_STYLE_NAMES["references"])
        _append_runs(p, parse_runs(cleaned))


def _suppress_numbering(paragraph: Paragraph) -> None:
    """Override the paragraph's numbering to numId=0 (no number) without
    changing the paragraph style."""
    pPr = paragraph._p.get_or_add_pPr()
    existing = pPr.find(qn("w:numPr"))
    if existing is not None:
        pPr.remove(existing)
    numPr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    numPr.append(ilvl)
    numId = OxmlElement("w:numId")
    numId.set(qn("w:val"), "0")
    numPr.append(numId)
    pPr.append(numPr)


# --- Shared run helpers --------------------------------------------------


def _append_runs(paragraph: Paragraph, runs: Iterable[Run], default_italic: bool = False) -> None:
    for run in runs:
        r = paragraph.add_run(run.text)
        if run.bold:
            r.bold = True
        if run.italic or (default_italic and not run.code):
            r.italic = True
        if run.code:
            r.font.name = "Courier New"


def _replace_paragraph_text(paragraph: Paragraph, runs: list[Run]) -> None:
    for r in list(paragraph.runs):
        r._element.getparent().remove(r._element)
    # Remove any stray empty w:r elements missed by paragraph.runs accessor
    for child in list(paragraph._p):
        if child.tag == qn("w:r"):
            paragraph._p.remove(child)
    _append_runs(paragraph, runs)


def _style_id_of(paragraph: Paragraph) -> str | None:
    return _style_id_of_elem(paragraph._p)


def _style_id_of_elem(p_elem) -> str | None:
    pPr = p_elem.find(qn("w:pPr"))
    if pPr is None:
        return None
    pStyle = pPr.find(qn("w:pStyle"))
    if pStyle is None:
        return None
    return pStyle.get(qn("w:val"))
